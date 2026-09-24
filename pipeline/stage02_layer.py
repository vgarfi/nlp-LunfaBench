"""Stage 02: layer.

Proposes the layer of each candidate from the regional marks of the two
lexical sources: kaikki (English Wiktionary) and Wikcionario (Spanish). Each
sense of the term is placed in a region (see regions.sense_region) and the
layer depends on which regions show up in either source:

- falso_amigo: there is a sense of a competing variety (Spain, Mexico) that is
  not used in the target variety and has a different gloss;
- polisemico_interno: no competitors, but other senses used in Buenos Aires
  (general or from broad regions) with a different gloss;
- lunfardo_sin_homonimo: only senses of the target variety (and at most of
  other regions that are not competitors);
- sin_clasificar: the term is in neither source.

The rioplatense sense comes from the lexical sources or, if they lack it, from
the collaborative dictionary (every term of diccionarioargentino.com is used in
Argentina). In that case glosses cannot be compared: the dictionary does not
mark regions and its definitions do not follow the Wiktionary format.

Outputs:
- data/interim/capas.jsonl: proposed layer, reason and senses by region.
- data/interim/02_capa_numeros.json: counts by layer and marks.

The agreement between the filter and the group is measured on the 90 curated
terms (stage 04), which uses cohen_kappa and read_annotations from here.

Usage:
    python -m pipeline.stage02_layer
"""

import argparse
import csv
import json
import os
import re
from collections import Counter, defaultdict

from pipeline import common, regions
from pipeline import source_kaikki as kk
from pipeline import source_wikcionario as wk

# Lexical sources: name, module and block of regional marks in the config.
LEXICAL_SOURCES = [("kaikki", kk, "variedades"),
                   ("wikcionario", wk, "variedades_wikcionario")]

INPUT = os.path.join(common.DATA_INTERIM, "candidatos.jsonl")
OUTPUT = os.path.join(common.DATA_INTERIM, "capas.jsonl")
NUMBERS = os.path.join(common.DATA_INTERIM, "02_capa_numeros.json")
ANNOTATION_DIR = os.path.join(common.REPO_ROOT, "data", "annotation")

LAYERS = ["falso_amigo", "polisemico_interno", "lunfardo_sin_homonimo", "sin_clasificar"]


def same_gloss(a, b, threshold):
    return common.same_gloss(a, b, threshold)


def senses_by_region(words, index, varieties, cfg):
    """Senses of a lexical source for the words of the candidate, grouped by
    region. Leaves out inflected forms, spelling variants and excluded parts
    of speech (proper names, etc.). Obsolete or dated senses do not count,
    except those of the target variety: dated lunfardo (mina = prostitute) is
    exactly the one of tango."""
    ignored = set(cfg["etapa02_capa"]["tags_ignorados"])
    excluded_pos = set(cfg["etapa01_candidatos"]["pos_excluidas"])
    by_region = defaultdict(list)
    for p in words:
        for e in index.get(p, []):
            if e["pos"] in excluded_pos:
                continue
            for a in e["acepciones"]:
                if a["forma_de"] or a["alt_de"]:
                    continue
                region, marks = regions.sense_region(a, varieties)
                if region != "objetivo" and ignored & set(a["tags"]):
                    continue
                by_region[region].append({
                    "id": a["id"], "palabra": p, "pos": e["pos"], "glosas": a["glosas"],
                    "marcas": marks,
                    # facha = looks: marked Rioplatense and also Spain
                    "tambien_competidora": region == "objetivo" and bool(
                        regions.group_marks(a, varieties["competidoras"])),
                })
    return by_region


def propose_layer(cand, sources, cfg):
    """Proposed layer from the evidence of the two lexical sources. Each
    source is looked at separately (kaikki glosses are in English and
    Wikcionario ones in Spanish, so they are only compared within the same
    source) and it is enough for one of them to show the competitor or the
    polysemy."""
    threshold = cfg["etapa02_capa"]["similitud_glosa"]
    by_source = {n: senses_by_region(cand[n]["palabras"], f["indice"], f["variedades"], cfg)
                 for n, f in sources.items()}
    marks = []
    if not any(any(pr.values()) for pr in by_source.values()):
        return "sin_clasificar", "no esta en kaikki ni en Wikcionario", by_source, marks, []
    if not any(pr["objetivo"] for pr in by_source.values()):
        marks.append("objetivo_solo_diccionario")
    if any(a["tambien_competidora"] for pr in by_source.values() for a in pr["objetivo"]):
        marks.append("objetivo_compartida_con_competidora")

    def differs(a, reference):
        return not any(same_gloss(a["glosas"], r["glosas"], threshold) for r in reference)

    competitors, others_here = {}, {}
    for n, pr in by_source.items():
        competitors[n] = [a for a in pr["competidora"]
                          if differs(a, pr["objetivo"] + pr["amplia"])]
        if pr["objetivo"]:
            others_here[n] = [a for a in pr["general"] + pr["amplia"] if differs(a, pr["objetivo"])]
        else:
            # If the rioplatense sense is in the other source (or another
            # language), there is no way to tell whether a general sense is
            # the same one
            distinct = []
            for a in pr["general"] + pr["amplia"]:
                if differs(a, distinct):
                    distinct.append(a)
            others_here[n] = distinct if len(distinct) >= 2 else []

    def summary(senses_by_source):
        parts = []
        for n, senses in senses_by_source.items():
            if senses:
                marks_ = {m for a in senses for m in a["marcas"]}
                # "Mexico" and "Mexican Spanish" say the same: show the labels
                regions_ = sorted({m for m in marks_ if not m.endswith(" Spanish")
                                   and not m.startswith("ES:")} or marks_)
                glosses = "; ".join(a["glosas"][-1] for a in senses[:2] if a["glosas"])
                parts.append(f"{n}" + (f" ({', '.join(regions_)})" if regions_ else "")
                             + f": {glosses}")
        return " | ".join(parts)

    if any(competitors.values()):
        deciding = [n for n in sources if competitors[n]]
        return "falso_amigo", summary(competitors), by_source, marks, deciding
    if any(others_here.values()):
        deciding = [n for n in sources if others_here[n]]
        return "polisemico_interno", summary(others_here), by_source, marks, deciding
    if any(pr["otra"] for pr in by_source.values()):
        marks.append("homonimo_en_otras_regiones")
    deciding = [n for n in sources if any(by_source[n].values())]
    return "lunfardo_sin_homonimo", "solo acepciones de la variedad objetivo", by_source, \
        marks, deciding


def classify(cfg):
    sources = {}
    for name, module, cfg_key in LEXICAL_SOURCES:
        sources[name] = {"indice": module.index(cfg), "variedades": cfg[cfg_key]}
    rows = []
    for cand in common.read_jsonl(INPUT):
        layer, reason, by_source, marks, deciding = propose_layer(cand, sources, cfg)
        rows.append({
            "term": cand["term"],
            "clave": cand["clave"],
            "fuentes": cand["fuentes"],
            "capa_propuesta": layer,
            "motivo": reason,
            "fuentes_que_deciden": deciding,
            "marcas": marks + cand.get("marcas", []),
            **{f"acepciones_{n}": {r: pr.get(r, []) for r in
                                   ("objetivo", "amplia", "competidora", "general", "otra")}
               for n, pr in by_source.items()},
            "pipeline": {"stage": "02_capa", "fecha": common.now_iso(), "validado": False},
        })
    common.write_jsonl(OUTPUT, rows)

    by_layer = Counter(f["capa_propuesta"] for f in rows)
    numbers = {
        "fecha": common.now_iso(),
        "parametros": {"variedades": cfg["variedades"],
                       "variedades_wikcionario": cfg["variedades_wikcionario"],
                       **cfg["etapa02_capa"]},
        "candidatos": len(rows),
        "por_capa": {c: by_layer.get(c, 0) for c in LAYERS},
        "por_capa_y_fuente_del_candidato": {
            c: dict(Counter("+".join(f["fuentes"]) for f in rows if f["capa_propuesta"] == c))
            for c in LAYERS},
        "por_capa_y_fuente_que_decide": {
            c: dict(Counter("+".join(f["fuentes_que_deciden"]) for f in rows
                            if f["capa_propuesta"] == c))
            for c in LAYERS},
        "marcas": dict(Counter(m for f in rows for m in f["marcas"])),
    }
    with open(NUMBERS, "w", encoding="utf-8") as f:
        json.dump(numbers, f, ensure_ascii=False, indent=2)
    print(json.dumps(numbers["por_capa"], ensure_ascii=False))
    print(f"-> {OUTPUT}")
    return rows


def cohen_kappa(pairs, labels):
    n = len(pairs)
    if not n:
        return None
    po = sum(a == b for a, b in pairs) / n
    pa = Counter(a for a, _ in pairs)
    pb = Counter(b for _, b in pairs)
    pe = sum(pa[e] * pb[e] for e in labels) / (n * n)
    return None if pe == 1 else (po - pe) / (1 - pe)


def read_annotations(path):
    """Annotated rows as dicts, from the CSV or from the .xlsx sheet (its
    headers are normalized: "capa anotador 1" -> capa_anotador_1)."""
    if not path.endswith(".xlsx"):
        with open(path, encoding="utf-8") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook
    sheet = load_workbook(path, data_only=True)["Anotación"]
    rows = sheet.iter_rows(values_only=True)
    headers = [re.sub(r"[^a-z0-9]+", "_", common.search_key(str(h or ""))).strip("_")
               for h in next(rows)]
    return [{k: ("" if v is None else str(v).strip()) for k, v in zip(headers, row)}
            for row in rows if row and row[0]]


def main():
    argparse.ArgumentParser(description=__doc__.split("\n")[0]).parse_args()
    classify(common.load_config())


if __name__ == "__main__":
    main()
