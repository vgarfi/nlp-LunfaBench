"""Stage 04: curated set.

The curated set has 90 items (40 false friends, 20 without a homonym, 15
polysemous and 15 controls, as in the report), only with lyric contexts and
full human validation.

Steps:
--select: draws exactly the quota of each layer among the terms with lyric
  fragments, with a fixed seed, prioritizing those with fragments from two
  different songs. Each term carries two fragments (A and B): the group picks
  the one that uses the sense, so no item is dropped. If no fragment of a
  term contains the term itself but another word with the same form (temo is
  from temer and not from temar; papá is not papa), the term is replaced by
  the next one of its layer in the draw.
  Outputs:
  - data/interim/curado_seleccion.jsonl: the 90 terms and their fragments
    (python -m pipeline.freeze copies them to data/curated/ with their
    provenance).
  - data/annotation/curado_planilla.xlsx: to share with the group.
  - data/annotation/curado_clave.csv: layer proposed by the filter. NOT to be
    shared: the annotation is blind.
--evaluate SHEET: Table 2 (proposed layer against the group's), agreement
  between annotators and the 90 items with their final layer and fragment:
  data/curated/curado_anotado.jsonl and reports/04_curado_numeros.json. It
  reads the frozen corpus in data/curated/.

Usage:
    python -m pipeline.stage04_curated --select
    python -m pipeline.freeze
    python -m pipeline.stage04_curated --evaluate data/annotation/curado_planilla.xlsx
"""

import argparse
import csv
import json
import math
import os
import random
import re
from collections import Counter, defaultdict

from wordfreq import top_n_list, zipf_frequency

from pipeline import common
from pipeline import source_kaikki
from pipeline import stage01_candidates as s01
from pipeline import stage02_layer as s02
from pipeline import stage02b_controls as s02b
from pipeline import stage03b_contexts as s03b

ANNOTATION_DIR = s02.ANNOTATION_DIR
SHEET = os.path.join(ANNOTATION_DIR, "curado_planilla.xlsx")
KEY_FILE = os.path.join(ANNOTATION_DIR, "curado_clave.csv")
SELECTION = os.path.join(common.DATA_INTERIM, "curado_seleccion.jsonl")
# The evaluation reads the frozen corpus (python -m pipeline.freeze), so it
# does not depend on the intermediate outputs.
FROZEN = os.path.join(common.REPO_ROOT, "data", "curated", "curado.jsonl")
ANNOTATED = os.path.join(common.REPO_ROOT, "data", "curated", "curado_anotado.jsonl")
NUMBERS = os.path.join(common.REPORTS, "04_curado_numeros.json")
SELECTION_NUMBERS = os.path.join(common.DATA_INTERIM, "04_seleccion_numeros.json")

# Layer of the benchmark and how the group annotates it (a control is
# general Spanish).
LAYERS = ["falso_amigo", "lunfardo_sin_homonimo", "polisemico_interno", "control_negativo"]
HUMAN_LAYER = {"falso_amigo": "falso_amigo", "lunfardo_sin_homonimo": "lunfardo_sin_homonimo",
               "polisemico_interno": "polisemico_interno", "control_negativo": "espanol_general"}
LAYER_OPTIONS = ["falso_amigo", "polisemico_interno", "lunfardo_sin_homonimo",
                 "espanol_general"]


# --- Selection -------------------------------------------------------------------

_ONLY_ANNOUNCES = re.compile(r"(?i)^\s*(dos|tres|varios|distintos)\s+"
                             r"(sentidos|significados|usos|acepciones)\b")


def reference_sense(target, candidates, controls):
    """The sense the fragment should use: the rioplatense one for a candidate
    (collaborative dictionary, else Wikcionario, else kaikki) and the general
    one for a control."""
    kind, key = target.split(":", 1)
    if kind == "ctrl":
        c = controls[key]
        return c["glosa"], "Wikcionario" if c["fuente_glosa"] == "wikcionario" else "Wiktionary"
    c = candidates[key]
    dic = c.get("diccionario_argentino")
    # "Dos sentidos" (chamuyo) announces senses but states none.
    definitions = [d for d in (dic or {}).get("definiciones", [])
                   if not _ONLY_ANNOUNCES.match(d["texto"])]
    if definitions:
        return definitions[0]["texto"], definitions[0]["url"]
    for name, label in (("wikcionario", "Wikcionario"), ("kaikki", "Wiktionary (en inglés)")):
        target_senses = c[name]["acepciones_objetivo"]
        if target_senses:
            return "; ".join(target_senses[0]["glosas"]), label
    return "", ""


def decade(x):
    year = x["source"].get("year") or x["source"].get("genius_release_year")
    return year // 10 * 10 if year else None


def distinct_songs(fragments):
    return len({x["source"]["genius_song_id"] for x in fragments})


def choose_fragments(fragments, used_by_decade, rng):
    """Fragments A and B of a term, from different songs if possible.
    Preferred: no guests from elsewhere, with a year (and from the decades
    least covered so far, to spread the curated set across eras) and
    longer."""
    rng.shuffle(fragments)
    order = sorted(fragments, key=lambda x: (
        not x["source"].get("invitados_de_ba", True),
        decade(x) is None,
        used_by_decade[decade(x)],
        -len(x["context"].split())))
    primary = order[0]
    alternative = next((x for x in order[1:]
                        if x["source"]["genius_song_id"] != primary["source"]["genius_song_id"]),
                       order[1] if len(order) > 1 else None)
    used_by_decade[decade(primary)] += 1
    return primary, alternative


def kaikki_readings(forms):
    """For each lyric form (normalized, with accents), the words it can be
    according to Wiktionary (kaikki): the lemma it inflects or is a variant
    of, or the form itself if it has senses of its own. If the form is not
    there as is (Genius lyrics lose accents), whatever has the same
    accent-free key counts. Proper names do not count: wordfreq lowercases
    everything."""
    keys = {common.search_key(f) for f in forms}
    by_word, by_key = defaultdict(set), defaultdict(set)
    for e in common.read_jsonl(source_kaikki.INDEX):
        key = common.search_key(e["word"])
        if key not in keys or e.get("pos") == "name":
            continue
        word = common.normalize(e["word"])
        for a in e["acepciones"]:
            lemmas = a["forma_de"] or a["alt_de"]
            readings = {common.normalize(l) for l in lemmas} if lemmas else {word}
            by_word[word] |= readings
            by_key[key] |= readings
    return {f: by_word.get(f) or by_key.get(common.search_key(f), set())
            for f in forms}


def is_other_word(fragment, own, term_zipf, readings):
    """The fragment does not contain the term but another word with the same
    form. If the lyric has an accent that belongs to no form of the term, the
    accent is enough: papá and papás are not papa, compás is not compa.
    Otherwise it counts when the other word is more frequent than the term:
    temo is from temer and not from temar, novio is a noun and not noviar.
    (wordfreq does not help with accents: it counts the papá written without
    an accent as papa.)"""
    form = common.normalize(fragment["forma_en_letra"])
    if form in own:
        return False
    all_readings = readings.get(form, set())
    others = all_readings - own
    if not others:
        return False
    accented = form != common.strip_accents(form)
    if (accented or common.strip_accents(form) in {common.strip_accents(p) for p in own}) \
            and not all_readings & own:
        return True
    return any(zipf_frequency(o, "es") > term_zipf for o in others)


def select_terms(cfg):
    ccfg = cfg["etapa04_curado"]
    rng = random.Random(ccfg["semilla"])
    frequent = {common.search_key(w) for w in top_n_list("es", ccfg["excluir_frecuentes"])}
    layers = {c["clave"]: c for c in common.read_jsonl(s02.OUTPUT)}
    candidates = {c["clave"]: c for c in common.read_jsonl(s01.OUTPUT)}
    controls = {c["clave"]: c for c in common.read_jsonl(s02b.OUTPUT)}

    fragments = defaultdict(list)
    for x in common.read_jsonl(s03b.CONTEXTS):
        if len(x["context"].replace(" / ", " ").split()) >= ccfg["min_palabras_contexto"]:
            fragments[x["objetivo"]].append(x)

    eligible = defaultdict(list)
    for target in sorted(fragments):
        kind, key = target.split(":", 1)
        if key in frequent or len(key) < 3:
            continue  # no, dar, hacer: candidates because of a slang use, but bad items
        if kind == "ctrl":
            eligible["control_negativo"].append(target)
            continue
        c = layers[key]
        if "posible_nombre_propio" in c["marcas"]:
            continue
        if c["capa_propuesta"] == "falso_amigo" and \
                "objetivo_compartida_con_competidora" in c["marcas"]:
            continue  # facha = looks is also used in Spain: not a clean false friend
        if c["capa_propuesta"] in LAYERS:
            eligible[c["capa_propuesta"]].append(target)

    # Exactly the quota of each layer, first among those with fragments from
    # two different songs (so there is always a B to choose).
    chosen, order = {}, {}
    for layer in LAYERS:
        with_two = [o for o in eligible[layer] if distinct_songs(fragments[o]) >= 2]
        with_one = [o for o in eligible[layer] if o not in with_two]
        rng.shuffle(with_two)
        rng.shuffle(with_one)
        order[layer] = with_two + with_one
        chosen[layer] = order[layer][:ccfg["cuotas"][layer]]

    def row_for(target, layer, fragments_, used_by_decade, rng_):
        primary, alternative = choose_fragments(list(fragments_), used_by_decade, rng_)
        text, source = reference_sense(target, candidates, controls)
        kind, key = target.split(":", 1)
        return {
            "objetivo": target, "term": primary["term"], "capa_propuesta": layer,
            "acepcion_de_referencia": text, "fuente_acepcion": source,
            "principal": primary, "alternativo": alternative,
            "emparejado_con": [e["falso_amigo"] for e in controls[key]["emparejado_con"]]
            if kind == "ctrl" else [],
        }

    rows, used_by_decade = [], Counter()
    for layer in LAYERS:
        for target in chosen[layer]:
            rows.append(row_for(target, layer, fragments[target], used_by_decade, rng))
    rng.shuffle(rows)  # the layers end up mixed in the sheet
    for i, row in enumerate(rows, 1):
        row["id"] = f"CUR-{i:03d}"

    # Replacements: if neither A nor B contains the term (only another word
    # with the same form), they are chosen again among its valid fragments,
    # and if it has none, the next term of its layer in the order of the draw
    # goes in. They have their own generator, so the rest of the draw stays
    # the same.
    readings = kaikki_readings({common.normalize(x["forma_en_letra"])
                                for layer in LAYERS for o in order[layer]
                                for x in fragments[o]})
    valid_cache = {}

    def valid_fragments(target, fragments_=None):
        kind, key = target.split(":", 1)
        own = {key}
        if kind == "cand":
            c = candidates[key]
            own |= {common.normalize(f) for f in c["formas"]} | \
                {v["clave"] for v in c["variantes"]}
        term = fragments[target][0]["term"]
        own.add(common.normalize(term))
        zipf = zipf_frequency(term, "es")
        if fragments_ is not None:
            return [x for x in fragments_ if not is_other_word(x, own, zipf, readings)]
        if target not in valid_cache:
            valid_cache[target] = valid_fragments(target, fragments[target])
        return valid_cache[target]

    replacement_rng = random.Random(f'{ccfg["semilla"]}-reemplazos')
    replacements = []
    for i, row in enumerate(rows):
        chosen_ab = [x for x in (row["principal"], row["alternativo"]) if x]
        if valid_fragments(row["objetivo"], chosen_ab):
            continue
        layer, used = row["capa_propuesta"], {f["objetivo"] for f in rows}
        if valid_fragments(row["objetivo"]):
            target = row["objetivo"]
        else:
            # First those with valid fragments from two different songs.
            target = sorted((o for o in order[layer] if o not in used and valid_fragments(o)),
                            key=lambda o: distinct_songs(valid_fragments(o)) < 2)[0]
        used_by_decade[decade(row["principal"])] -= 1
        new_row = row_for(target, layer, valid_fragments(target), used_by_decade, replacement_rng)
        new_row["id"] = row["id"]
        rows[i] = new_row
        replacements.append({"id": row["id"], "capa": layer, "sale": row["term"],
                             "forma_en_letra": sorted({x["forma_en_letra"].lower()
                                                       for x in chosen_ab}),
                             "entra": new_row["term"]})
    # The fragment that carries another word stays marked (cocina for
    # cocinar): the group picks the other one.
    for row in rows:
        ab = [x for x in (row["principal"], row["alternativo"]) if x]
        good_ids = [id(x) for x in valid_fragments(row["objetivo"], ab)]
        for x in ab:
            x["forma_de_otra_palabra"] = id(x) not in good_ids

    common.write_jsonl(SELECTION, rows)
    with open(KEY_FILE, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "term", "capa_propuesta", "objetivo", "semilla"])
        for row in rows:
            w.writerow([row["id"], row["term"], row["capa_propuesta"], row["objetivo"],
                        ccfg["semilla"]])
    build_sheet(rows, cfg)
    numbers = {
        "fecha": common.now_iso(), "semilla": ccfg["semilla"],
        "elegibles": {k: len(v) for k, v in eligible.items()},
        "seleccionados": dict(Counter(f["capa_propuesta"] for f in rows)),
        "reemplazados": replacements,
        "decadas_de_los_fragmentos": {str(k): v for k, v in
                                      sorted(used_by_decade.items(),
                                             key=lambda x: (x[0] is None, x[0] or 0)) if v},
    }
    with open(SELECTION_NUMBERS, "w", encoding="utf-8") as f:
        json.dump(numbers, f, ensure_ascii=False, indent=2)
    print(json.dumps(numbers, ensure_ascii=False, indent=2))
    print(f"-> {SHEET}")


# --- Annotation sheet (its texts are in Spanish: it is for the group) ------------

COLUMNS = [  # (header, width, filled in by hand)
    ("id", 9, False), ("término", 13, False), ("acepción de referencia", 42, False),
    ("fuente", 16, False),
    ("fragmento A", 46, False), ("canción A", 24, False),
    ("fragmento B", 46, False), ("canción B", 24, False),
    ("anotador 1", 11, True), ("capa 1", 20, True), ("fragmento 1", 11, True),
    ("acepción competidora 1", 24, True),
    ("anotador 2", 11, True), ("capa 2", 20, True), ("fragmento 2", 11, True),
    ("acepción competidora 2", 24, True),
    ("coincide capa", 9, False), ("coincide fragmento", 10, False),
    ("capa final", 20, True), ("fragmento final", 11, True), ("comentario", 30, True),
]

INSTRUCTIONS = [
    ("LunfaBench · Planilla del conjunto curado", "titulo"),
    ("Son los 90 términos del conjunto curado: 40 falsos amigos, 20 sin homónimo, 15 "
     "polisémicos y 15 controles, mezclados. Cada uno trae dos fragmentos reales (A y B) de "
     "canciones de artistas de Buenos Aires.", None),
    ("Qué hace cada anotador", "subtitulo"),
    ("1. Escribí tu nombre en «anotador». Cada fila la anotan dos personas, por separado: no "
     "mires la columna de la otra persona hasta terminar.", None),
    ("2. En «capa», elegí de la lista qué tipo de palabra es, pensando en Buenos Aires:", None),
    ("   falso_amigo: tiene una acepción de España o México que acá no se usa (curro: acá "
     "negocio turbio, en España trabajo).", None),
    ("   polisemico_interno: tiene varias acepciones y todas se usan acá (pileta: piscina y "
     "pileta de la cocina).", None),
    ("   lunfardo_sin_homonimo: no existe con otro sentido en España ni en México (bondi, "
     "chamuyo).", None),
    ("   espanol_general: es una palabra común del español, que significa lo mismo en todos "
     "lados (ventana).", None),
    ("3. En «fragmento», elegí A o B: el que usa la acepción de referencia. Si los dos la "
     "usan, el que mejor se entienda. Si ninguno, el más cercano, y contalo en «comentario».",
     None),
    ("4. Si elegiste falso_amigo, escribí la acepción competidora y dónde se usa, por ejemplo "
     "«trabajo (España)».", None),
    ("Se puede consultar el DLE y el Diccionario de americanismos. No consultes Wiktionary ni "
     "Wikcionario: son la fuente del filtro automático y el acuerdo quedaría inflado.", None),
    ("Después: decisión final", "subtitulo"),
    ("Cuando los dos anotadores terminan, las columnas «coincide» se completan solas. Si "
     "coinciden, copien la respuesta en «capa final» y «fragmento final»; si no, "
     "discútanla y dejen el motivo en «comentario».", None),
    ("Celdas a completar", "subtitulo"),
    ("Solo las celdas amarillas de la hoja Anotación. Las demás vienen del pipeline o son "
     "fórmulas. La hoja Resumen muestra el avance y el acuerdo.", None),
    ("Ejemplo de fila completa (el término no está en la planilla)", "subtitulo"),
]


def fragment_origin(x):
    year = x["source"].get("genius_release_year")
    return f'{x["source"]["song"]} · {x["source"]["artist"]}' + (f" ({year})" if year else "")


def build_sheet(rows, cfg):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    normal_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    gray = PatternFill("solid", fgColor="E7E6E1")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    border = Border(bottom=Side(style="thin", color="D0CFC9"))
    wrap = Alignment(wrap_text=True, vertical="top")
    col = {h: get_column_letter(i) for i, (h, _, _) in enumerate(COLUMNS, 1)}

    workbook = Workbook()
    inst = workbook.active
    inst.title = "Instrucciones"
    inst.column_dimensions["A"].width = 120
    row_num = 1
    for text, style in INSTRUCTIONS:
        cell = inst.cell(row=row_num, column=1, value=text)
        cell.font = Font(name="Arial", size=14 if style == "titulo" else 10,
                         bold=style in ("titulo", "subtitulo"))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_num += 2 if style == "titulo" else 1
    # Example of a full row, with a term that is not in the sheet.
    used = {f["term"] for f in rows}
    example_term = next(t for t in ("curro", "boliche", "torta", "coger", "concha", "cajeta")
                        if t not in used)
    example = {"término": example_term, "anotador 1": "Tomás", "capa 1": "falso_amigo",
               "fragmento 1": "A", "acepción competidora 1": "trabajo, empleo (España)"
               if example_term == "curro" else "(la acepción de España o México)"}
    row_num += 1
    example_headers = [h for h, _, _ in COLUMNS if h in example or h == "acepción de referencia"]
    for j, h in enumerate(example_headers, 1):
        inst.cell(row=row_num, column=j, value=h).font = bold_font
        value = example.get(h, "Negocio turbio o arreglo deshonesto"
                            if example_term == "curro" else "")
        inst.cell(row=row_num + 1, column=j, value=value).font = normal_font

    sheet = workbook.create_sheet("Anotación")
    for j, (h, width, _) in enumerate(COLUMNS, 1):
        c = sheet.cell(row=1, column=j, value=h)
        c.font, c.fill, c.alignment = bold_font, gray, Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(j)].width = width
    only_a = []
    for i, f in enumerate(rows, 2):
        a, b = f["principal"], f["alternativo"]
        if not b:
            only_a.append(i)
        cell_values = {
            "id": f["id"], "término": f["term"],
            "acepción de referencia": f["acepcion_de_referencia"],
            "fuente": f["fuente_acepcion"],
            "fragmento A": a["context"], "canción A": fragment_origin(a),
            "fragmento B": b["context"] if b else "", "canción B": fragment_origin(b) if b else "",
            "coincide capa": f'=IF(OR({col["capa 1"]}{i}="",{col["capa 2"]}{i}=""),"",'
                             f'IF({col["capa 1"]}{i}={col["capa 2"]}{i},"sí","no"))',
            "coincide fragmento": f'=IF(OR({col["fragmento 1"]}{i}="",{col["fragmento 2"]}{i}=""),'
                                  f'"",IF({col["fragmento 1"]}{i}={col["fragmento 2"]}{i},'
                                  f'"sí","no"))',
        }
        for j, (h, _, by_hand) in enumerate(COLUMNS, 1):
            c = sheet.cell(row=i, column=j, value=cell_values.get(h, ""))
            c.font, c.alignment, c.border = normal_font, wrap, border
            if by_hand:
                c.fill = yellow
            if h == "fuente" and str(cell_values["fuente"]).startswith("http"):
                c.hyperlink, c.value = cell_values["fuente"], "diccionarioargentino.com"
                c.font = Font(name="Arial", size=10, color="1F5FAD", underline="single")
        length = max(len(f["acepcion_de_referencia"]) / 40, len(a["context"]) / 44,
                     len(b["context"]) / 44 if b else 0, 1)
        sheet.row_dimensions[i].height = min(15 * math.ceil(length), 120)
    last_row = len(rows) + 1
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    dv_layer = DataValidation(type="list", formula1='"' + ",".join(LAYER_OPTIONS) + '"',
                              allow_blank=True, showErrorMessage=True,
                              error="Elegí una opción de la lista")
    sheet.add_data_validation(dv_layer)
    for h in ("capa 1", "capa 2", "capa final"):
        dv_layer.add(f"{col[h]}2:{col[h]}{last_row}")
    # A or B; in the rows without fragment B, only A.
    dv_ab = DataValidation(type="list", formula1='"A,B"', allow_blank=True,
                           showErrorMessage=True, error="Elegí A o B")
    dv_a = DataValidation(type="list", formula1='"A"', allow_blank=True,
                          showErrorMessage=True, error="Esta fila solo tiene fragmento A")
    sheet.add_data_validation(dv_ab)
    sheet.add_data_validation(dv_a)
    for h in ("fragmento 1", "fragmento 2", "fragmento final"):
        for i in range(2, last_row + 1):
            (dv_a if i in only_a else dv_ab).add(f"{col[h]}{i}")
    red = PatternFill("solid", fgColor="F8D7D3")
    for h in ("coincide capa", "coincide fragmento"):
        sheet.conditional_formatting.add(f"{col[h]}2:{col[h]}{last_row}",
                                         CellIsRule(operator="equal", formula=['"no"'], fill=red))

    res = workbook.create_sheet("Resumen")
    res.column_dimensions["A"].width = 34
    res.column_dimensions["B"].width = 14
    col_range = lambda h: f"'Anotación'!${col[h]}$2:${col[h]}${last_row}"  # noqa: E731
    summary_rows = [
        ("Avance", None),
        ("Términos", f'=COUNTIF({col_range("id")},"?*")'),
        ("Con anotador 1", f'=COUNTIF({col_range("capa 1")},"?*")'),
        ("Con anotador 2", f'=COUNTIF({col_range("capa 2")},"?*")'),
        ("Con capa final", f'=COUNTIF({col_range("capa final")},"?*")'),
        ("Acuerdo en capa", f'=IFERROR(COUNTIF({col_range("coincide capa")},"sí")/'
                            f'(COUNTIF({col_range("coincide capa")},"sí")+'
                            f'COUNTIF({col_range("coincide capa")},"no")),"")'),
        ("Acuerdo en fragmento", f'=IFERROR(COUNTIF({col_range("coincide fragmento")},"sí")/'
                                 f'(COUNTIF({col_range("coincide fragmento")},"sí")+'
                                 f'COUNTIF({col_range("coincide fragmento")},"no")),"")'),
        (None, None),
        ("Capa final", "Términos"),
    ] + [(layer, f'=COUNTIF({col_range("capa final")},"{layer}")') for layer in LAYER_OPTIONS]
    for r, (a, b) in enumerate(summary_rows, 1):
        if a is not None:
            res.cell(row=r, column=1, value=a).font = bold_font if b is None or r == 9 else normal_font
        if b is not None:
            res.cell(row=r, column=2, value=b).font = bold_font if r == 9 else normal_font
    for r in (6, 7):
        res.cell(row=r, column=2).number_format = "0%"
    os.makedirs(ANNOTATION_DIR, exist_ok=True)
    workbook.save(SHEET)


# --- Evaluation -------------------------------------------------------------------

def evaluate(path, cfg):
    annotated = {r["id"]: r for r in s02.read_annotations(path)}
    key_rows = {r["id"]: r for r in csv.DictReader(open(KEY_FILE, encoding="utf-8"))}
    corpus = {f["id"]: f for f in common.read_jsonl(FROZEN)}

    def final(r, field):
        """Final decision; if missing, the matching one or annotator 1's."""
        return (r.get(f"{field}_final") or r.get(f"{field}_1") or "").strip()

    proposed_vs_group, between_layer, between_fragment = [], [], []
    for id_, r in annotated.items():
        if final(r, "capa"):
            proposed_vs_group.append((HUMAN_LAYER[key_rows[id_]["capa_propuesta"]], final(r, "capa")))
        if r.get("capa_1") and r.get("capa_2"):
            between_layer.append((r["capa_1"], r["capa_2"]))
        if r.get("fragmento_1") and r.get("fragmento_2"):
            between_fragment.append((r["fragmento_1"], r["fragmento_2"]))

    result = {
        "anotados": len(proposed_vs_group),
        "tabla2_propuesta_vs_grupo": {p: {h: sum(1 for a, b in proposed_vs_group if a == p and b == h)
                                          for h in LAYER_OPTIONS}
                                      for p in sorted({a for a, _ in proposed_vs_group})},
        "tabla2_acuerdo": (sum(a == b for a, b in proposed_vs_group) / len(proposed_vs_group))
        if proposed_vs_group else None,
        "tabla2_kappa": s02.cohen_kappa(proposed_vs_group, LAYER_OPTIONS),
        "entre_anotadores_capa": {"pares": len(between_layer),
                                  "kappa": s02.cohen_kappa(between_layer, LAYER_OPTIONS)},
        "entre_anotadores_fragmento": {"pares": len(between_fragment),
                                       "kappa": s02.cohen_kappa(between_fragment, ["A", "B"])},
    }

    # The 90 items with the layer and the fragment the group decided.
    from_human = {v: k for k, v in HUMAN_LAYER.items()}
    items = []
    for n, id_ in enumerate(sorted(corpus), 1):
        f, r = corpus[id_], annotated.get(id_, {})
        layer = from_human.get(final(r, "capa"), f["capa_propuesta"])
        x = f["alternativo"] if final(r, "fragmento") == "B" and f["alternativo"] else f["principal"]
        items.append({
            "id": f"LB-{n:03d}", "term": f["term"], "layer": layer, "decade": decade(x),
            "genre": x.get("genre", ""), "context": x["context"], "context_source": "letra",
            "source": {**x["source"], "planilla_id": id_},
            "reference_sense": {"text": f["acepcion_de_referencia"],
                                "source": f["fuente_acepcion"]},
            "competing_sense": r.get("acepcion_competidora_1") or
            r.get("acepcion_competidora_2") or "",
            "options": {"a_correcta": "", "b_otra_variedad": "", "c_plausible": "",
                        "d_incorrecta": ""},
            "pipeline": {"stage": "04_curado", "validated_sense": bool(final(r, "fragmento")),
                         "validated_item": False, "capa_propuesta": f["capa_propuesta"]},
            "annotation": {"raw": {k: v for k, v in r.items()
                                   if k.startswith(("capa", "fragmento", "anotador"))},
                           "adjudicated": final(r, "capa"), "doubt": False},
        })
    common.write_jsonl(ANNOTATED, items)
    result["curado_por_capa"] = dict(Counter(x["layer"] for x in items))
    with open(NUMBERS, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    print(f"-> {ANNOTATED}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--select", action="store_true")
    ap.add_argument("--evaluate", metavar="SHEET")
    args = ap.parse_args()
    cfg = common.load_config()
    if args.select:
        select_terms(cfg)
    if args.evaluate:
        evaluate(args.evaluate, cfg)


if __name__ == "__main__":
    main()
